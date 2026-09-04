import io

from openpyxl import load_workbook


def parse_names(data: bytes) -> list:
    """Parse an uploaded Excel file and return the list of student names.

    Expects a header row containing a column named 'Name'.
    """
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ValueError("Please upload a valid Excel file.")

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError("The Excel file appears to be empty.")

    name_idx = None
    for i, cell in enumerate(header):
        if cell is not None and str(cell).strip().lower() == "name":
            name_idx = i
            break

    if name_idx is None:
        raise ValueError('Excel file must contain a column named "Name".')

    names = []
    for row in rows:
        if row is None:
            continue
        cell = row[name_idx] if name_idx < len(row) else None
        if cell is None:
            continue
        val = str(cell).strip()
        if val:
            names.append(val)

    if not names:
        raise ValueError("The Name column does not contain any student names.")

    return names
